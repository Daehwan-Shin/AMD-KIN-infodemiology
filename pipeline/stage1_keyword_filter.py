"""
Stage 1: Apply deterministic keyword filter to 20,797 rows.
Sample 50 items for pilot codebook v2 validation.

Output:
  _stage1_candidates.jsonl  : all rows passing Stage 1 filter
  _pilot_sample_50.jsonl    : 50 randomly sampled items (seed=42)
"""
import sys
import re
import json
import random
sys.stdout.reconfigure(encoding='utf-8')

import openpyxl

SRC = '황반변성_지식인_크롤링 260415(again).xlsx'

# Stage 1 Include patterns (must match in question_title OR question_content)
INCLUDE_PATTERNS = [
    re.compile(r'황반\s*변성', re.IGNORECASE),
    re.compile(r'\bAMD\b'),
    re.compile(r'macular\s+degeneration', re.IGNORECASE),
]

# Stage 1 Exclude patterns (computer-hardware noise)
EXCLUDE_PATTERNS = [
    re.compile(r'라이젠'),
    re.compile(r'라데온'),
    re.compile(r'\bRyzen\b', re.IGNORECASE),
    re.compile(r'\bRadeon\b', re.IGNORECASE),
    re.compile(r'그래픽\s*카드'),
    re.compile(r'\bGPU\b'),
    re.compile(r'\bCPU\b'),
]


def matches_any(text, patterns):
    if not text:
        return False
    return any(p.search(text) for p in patterns)


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb['Sheet1']
    header = [c.value for c in ws[1]]
    ix = {h: i for i, h in enumerate(header)}

    total = 0
    n_include_match = 0
    n_exclude_match = 0
    n_pass = 0
    candidates = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        total += 1
        title = row[ix['question_title']] or ''
        content = row[ix['question_content']] or ''
        text = f'{title}\n{content}'

        has_amd_kw = matches_any(text, INCLUDE_PATTERNS)
        if has_amd_kw:
            n_include_match += 1
        has_noise = matches_any(text, EXCLUDE_PATTERNS)
        if has_noise:
            n_exclude_match += 1

        if has_amd_kw and not has_noise:
            n_pass += 1
            candidates.append({
                'row_index': total + 1,  # 1-based excel row including header
                'url': row[ix['url']],
                'question_title': title,
                'question_content': content,
                'question_date': str(row[ix['question_date']] or ''),
                'category': row[ix['category']] or '',
                'question_views': row[ix['question_views']],
                'answer_count': row[ix['answer_count']],
            })

    print(f'Total rows: {total}')
    print(f'Has AMD keyword in question: {n_include_match}')
    print(f'Has hardware-noise keyword: {n_exclude_match}')
    print(f'Stage 1 pass (include AND NOT exclude): {n_pass}')

    with open('_stage1_candidates.jsonl', 'w', encoding='utf-8') as f:
        for c in candidates:
            f.write(json.dumps(c, ensure_ascii=False) + '\n')
    print(f'Saved: _stage1_candidates.jsonl  ({n_pass} items)')

    # Random sample 50 with fixed seed
    random.seed(42)
    pilot = random.sample(candidates, 50)
    # Assign pilot IDs
    for i, c in enumerate(pilot, 1):
        c['pilot_id'] = f'P-{i:03d}'

    with open('_pilot_sample_50.jsonl', 'w', encoding='utf-8') as f:
        for c in pilot:
            f.write(json.dumps(c, ensure_ascii=False) + '\n')
    print(f'Saved: _pilot_sample_50.jsonl  (50 items, seed=42)')


if __name__ == '__main__':
    main()
