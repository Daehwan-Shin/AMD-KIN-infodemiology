"""
Stage 4 주제 분류 머지 + 엄격 검증:
- 4 chunks 로드 → id→topic
- _final_strict_AMD.jsonl 2,048 id와 대조 (누락/잉여/중복)
- primary/secondary C1-C9 유효성
- 통과 시 최종 데이터셋 + 분포 + 엑셀
"""
import json, sys
sys.stdout.reconfigure(encoding='utf-8')
from collections import Counter

VALID = {f'C{i}' for i in range(1, 10)}
CNAME = {
 'C1':'질환정보·원인','C2':'증상·진단·검사','C3':'치료일반·수술','C4':'항VEGF 주사치료',
 'C5':'영양·생활습관·예방','C6':'경과·예후·실명우려','C7':'비용·보험·장애·행정',
 'C8':'병원·의료진','C9':'기타(부작용·상호작용)'}

# 원본 strict_AMD
with open('_final_strict_AMD.jsonl', 'r', encoding='utf-8') as f:
    base = [json.loads(l) for l in f if l.strip()]
base_ids = []
base_map = {}
for r in base:
    sid = r.get('phase2_id') or r.get('phase1_id')
    base_ids.append(sid)
    base_map[sid] = r
base_idset = set(base_ids)
print(f'원본 strict_AMD: {len(base)}건 (unique id {len(base_idset)})')

# chunks 로드
topic = {}
dups = []
bad = []
for i in range(1, 5):
    with open(f'_stage4_topic_chunk_{i}.jsonl', 'r', encoding='utf-8') as f:
        cnt = 0
        for ln in f:
            ln = ln.strip()
            if not ln:
                continue
            try:
                o = json.loads(ln)
            except json.JSONDecodeError:
                bad.append((i, ln[:80])); continue
            sid = o.get('id')
            pt = o.get('primary_topic')
            st = o.get('secondary_topics', [])
            if sid in topic:
                dups.append(sid)
            if pt not in VALID:
                bad.append((i, f'{sid} bad primary {pt}'))
            if not isinstance(st, list) or any(s not in VALID for s in st):
                bad.append((i, f'{sid} bad secondary {st}'))
            topic[sid] = {'primary': pt, 'secondary': [s for s in st if s in VALID],
                          'reason': o.get('reason', '')}
            cnt += 1
    print(f'  chunk {i}: {cnt}건')

print(f'\n수집된 topic id: {len(topic)}')
missing = [s for s in base_ids if s not in topic]      # 분류 안 된 원본
extra = [s for s in topic if s not in base_idset]       # 원본에 없는 잉여
print(f'누락(원본엔 있는데 분류 없음): {len(missing)}건')
print(f'잉여(분류엔 있는데 원본 없음): {len(extra)}건')
print(f'중복 id: {len(set(dups))}건 {sorted(set(dups))[:10]}')
print(f'형식 오류: {len(bad)}건')
for b in bad[:15]:
    print('   ', b)
if missing:
    print(f'\n⚠ 누락 id 목록(앞 30): {missing[:30]}')
    with open('_stage4_missing_ids.txt', 'w', encoding='utf-8') as f:
        for m in missing:
            f.write(m + '\n')
    print('   → _stage4_missing_ids.txt 저장 (재분류 대상)')
if extra:
    print(f'⚠ 잉여 id(앞 20): {extra[:20]}')

ok = (len(missing) == 0 and len(set(dups)) == 0 and len(bad) == 0)
print(f'\n검증 {"✅ 통과" if ok else "❌ 실패 — 누락/오류 해결 후 재머지 필요"}')

# 통과 시 최종 산출
if ok:
    final = []
    for sid in base_ids:
        r = dict(base_map[sid])
        t = topic[sid]
        r['primary_topic'] = t['primary']
        r['secondary_topics'] = t['secondary']
        r['topic_reason'] = t['reason']
        final.append(r)
    with open('_stage4_topic_final.jsonl', 'w', encoding='utf-8') as f:
        for r in final:
            f.write(json.dumps(r, ensure_ascii=False) + '\n')
    print(f'\n_stage4_topic_final.jsonl 저장: {len(final)}건')

    pri = Counter(r['primary_topic'] for r in final)
    print('\n=== Primary topic 분포 ===')
    for c in [f'C{i}' for i in range(1, 10)]:
        n = pri.get(c, 0)
        bar = '█' * round(n / len(final) * 60)
        print(f'  {c} {CNAME[c]:<16} {n:>4} ({n/len(final)*100:4.1f}%) {bar}')

    # primary+secondary 합산(주제 전체 노출 빈도)
    anyt = Counter()
    for r in final:
        anyt[r['primary_topic']] += 1
        for s in r['secondary_topics']:
            anyt[s] += 1
    print('\n=== Primary+Secondary 합산 (주제 언급 빈도) ===')
    for c in [f'C{i}' for i in range(1, 10)]:
        n = anyt.get(c, 0)
        print(f'  {c} {CNAME[c]:<16} {n:>4} ({n/len(final)*100:4.1f}%)')

    nsec = Counter(len(r['secondary_topics']) for r in final)
    print(f'\nsecondary 개수 분포: {dict(sorted(nsec.items()))}')
    multi = sum(1 for r in final if r['secondary_topics'])
    print(f'multi-label(secondary≥1): {multi}건 ({multi/len(final)*100:.1f}%)')

    # 엑셀
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'topic_classified'
    cols = ['No','id','source','primary_topic','primary_name','secondary_topics',
            'question_title','question_content','topic_reason',
            'codex','opus','gemini','url','category','question_date']
    ws.append(cols)
    for ci,h in enumerate(cols,1):
        c=ws.cell(row=1,column=ci); c.font=Font(bold=True,color='FFFFFF')
        c.fill=PatternFill('solid',fgColor='4F81BD'); c.alignment=Alignment(horizontal='center',vertical='center')
    for i,r in enumerate(final,1):
        sid=r.get('phase2_id') or r.get('phase1_id')
        ws.append([i,sid,r.get('source',''),r['primary_topic'],CNAME[r['primary_topic']],
                   ','.join(r['secondary_topics']),
                   r.get('question_title',''),r.get('question_content',''),r.get('topic_reason',''),
                   r.get('codex_AMD_status',''),r.get('opus_AMD_status',''),r.get('gemini_AMD_status',''),
                   r.get('url',''),r.get('category',''),r.get('question_date','')])
    for col,w in zip('ABCDEFGHIJKLMNO',[5,9,13,11,16,14,30,48,34,16,16,16,24,12,12]):
        ws.column_dimensions[col].width=w
    for row in ws.iter_rows(min_row=2,max_row=ws.max_row):
        for cell in row: cell.alignment=Alignment(wrap_text=True,vertical='top')
    ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:O{ws.max_row}'

    wss=wb.create_sheet('summary')
    wss.column_dimensions['A'].width=14; wss.column_dimensions['B'].width=22
    wss.column_dimensions['C'].width=10; wss.column_dimensions['D'].width=10
    wss.append(['code','category','primary','prim+sec']);
    for c in [f'C{i}' for i in range(1,10)]:
        wss.append([c,CNAME[c],pri.get(c,0),anyt.get(c,0)])
    wss.append(['','TOTAL',len(final),sum(anyt.values())])
    for cell in wss[1]: cell.font=Font(bold=True)

    wb.save('AMD_topic_classified_2048.xlsx')
    print('→ AMD_topic_classified_2048.xlsx (topic_classified + summary)')
