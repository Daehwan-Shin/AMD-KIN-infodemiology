# -*- coding: utf-8 -*-
"""v5b 최종 분석 (라벨링 완료 후):
1. v5b-balanced 100 sample (86 기존 + 14 추가) → SDH/NKT 라벨 통합
2. v5b rule 재평가: κ, gold-subset
3. disagreement 17 final consensus 통합 → gold-subset 확장
"""
import json, sys, re
sys.stdout.reconfigure(encoding='utf-8')
from collections import Counter
from sklearn.metrics import cohen_kappa_score, confusion_matrix
import openpyxl

# === load key (v5b balanced) ===
with open('_R1_validation_key_v5b_balanced.jsonl', encoding='utf-8') as f:
    key = [json.loads(l) for l in f if l.strip()]
key_by_id = {k['id']: k for k in key}
print(f'v5b-balanced key: {len(key)} cases (original + added)')

# === load original 100 labels (after SDH update) ===
def load(p):
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb['data']; rows = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        no, rid, _, _, lbl, _ = r
        rows[rid] = (lbl or '').strip()
    wb.close(); return rows
sdh_orig = load('AMD_R1_validation_100_SDH_Done.xlsx')
nkt_orig = load('AMD_R1_validation_100_NKT_Done.xlsx')
print(f'SDH original: {Counter(sdh_orig.values())}')
print(f'NKT original: {Counter(nkt_orig.values())}')

# === load ADD15 (new labels) ===
def load_add(p):
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb['data']; rows = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if len(r) < 6: continue
        no, rid, _, _, _, lbl, *_ = r
        rows[rid] = (lbl or '').strip()
    wb.close(); return rows
sdh_add = load_add('수정용자료/AMD_R1_validation_ADD15_v5b_SDH_done.xlsx')
nkt_add = load_add('수정용자료/AMD_R1_validation_ADD15_v5b_NKT__NKTdone.xlsx')
print(f'\nADD15 SDH: {Counter(sdh_add.values())}')
print(f'ADD15 NKT: {Counter(nkt_add.values())}')

# === merge: each rid → SDH label, NKT label ===
all_sdh = {**sdh_orig, **sdh_add}
all_nkt = {**nkt_orig, **nkt_add}

# === build v5b-balanced 100 sample data ===
records = []
for k in key:
    rid = k['id']
    sdh_lbl = all_sdh.get(rid, '')
    nkt_lbl = all_nkt.get(rid, '')
    rule_v5b = k.get('rule_label_v5b', '')
    records.append({
        'id': rid, 'sample_no': k['sample_no'],
        'origin': k.get('origin', '?'),
        'sdh': sdh_lbl, 'nkt': nkt_lbl, 'rule': rule_v5b,
    })

# 확인: 라벨 누락 케이스
missing_sdh = [r for r in records if r['sdh'] not in ('R1','not_R1','uncertain') or not r['sdh']]
missing_nkt = [r for r in records if r['nkt'] not in ('R1','not_R1','uncertain') or not r['nkt']]
print(f'\nLabel coverage:')
print(f'  SDH labeled: {len(records) - len(missing_sdh)} / {len(records)}')
print(f'  NKT labeled: {len(records) - len(missing_nkt)} / {len(records)}')
if missing_sdh:
    print(f'  SDH missing: {[r["id"] for r in missing_sdh[:5]]}')
if missing_nkt:
    print(f'  NKT missing: {[r["id"] for r in missing_nkt[:5]]}')

# === distributions ===
print(f'\n=== v5b-balanced 100 sample 분포 ===')
print(f'  rule_v5b:  {Counter(r["rule"] for r in records)}')
print(f'  SDH:       {Counter(r["sdh"] for r in records)}')
print(f'  NKT:       {Counter(r["nkt"] for r in records)}')

# === κ + raw + gold-subset ===
def kappa(rule_key, label_key, name):
    pairs = []
    for r in records:
        rl = r[rule_key]; el = r[label_key]
        if rl in ('R1','not_R1') and el in ('R1','not_R1'):
            pairs.append((rl, el))
    if not pairs:
        print(f'  {name}: no valid pairs'); return
    yr, ye = zip(*pairs)
    yt = [['R1','not_R1'].index(x) for x in yr]
    yp = [['R1','not_R1'].index(x) for x in ye]
    raw = sum(1 for a,b in zip(yt,yp) if a==b)/len(yt)
    k = cohen_kappa_score(yt, yp)
    cm = confusion_matrix(yr, ye, labels=['R1','not_R1'])
    print(f'  {name}  n={len(pairs)}  κ={k:.3f}  raw={raw*100:.1f}%')
    print(f'    rule_R1→exp_R1={cm[0,0]} rule_R1→exp_not={cm[0,1]} | rule_not→exp_R1={cm[1,0]} rule_not→exp_not={cm[1,1]}')

print(f'\n=== v5b rule vs expert ===')
kappa('rule', 'sdh', 'SDH vs v5b')
kappa('rule', 'nkt', 'NKT vs v5b')

print(f'\n=== inter-expert (SDH vs NKT) ===')
kappa('sdh', 'nkt', 'SDH vs NKT')

# === Gold-subset (양 expert 일치 + 양쪽 R1 또는 not_R1) ===
print(f'\n=== Gold-subset (SDH ∩ NKT) ===')
gold = [r for r in records if r['sdh'] == r['nkt'] and r['sdh'] in ('R1','not_R1')]
gold_match = sum(1 for r in gold if r['rule'] == r['sdh'])
gold_r1 = [r for r in gold if r['sdh'] == 'R1']
gold_not = [r for r in gold if r['sdh'] == 'not_R1']
gold_r1_m = sum(1 for r in gold_r1 if r['rule'] == 'R1')
gold_not_m = sum(1 for r in gold_not if r['rule'] == 'not_R1')
print(f'  gold n={len(gold)}, rule match={gold_match}/{len(gold)} = {100*gold_match/len(gold):.1f}%')
print(f'  gold-R1     {gold_r1_m}/{len(gold_r1)} = {100*gold_r1_m/len(gold_r1) if gold_r1 else 0:.1f}%')
print(f'  gold-not_R1 {gold_not_m}/{len(gold_not)} = {100*gold_not_m/len(gold_not) if gold_not else 0:.1f}%')

# === disagreement 17 final consensus 통합 ===
print(f'\n{"="*70}')
print(f'=== Disagreement 17 FINAL CONSENSUS 통합 ===')
print(f'{"="*70}')
wb = openpyxl.load_workbook('수정용자료/AMD_R1_inter_expert_disagreement_v2_NKTdone.xlsx',
                            read_only=True, data_only=True)
ws = wb['disagreement']
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
final_col = headers.index('★ FINAL LABEL (합의)')
consensus_data = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    rid = r[1]
    final_lbl = (r[final_col] or '').strip()
    if final_lbl: consensus_data[rid] = final_lbl
wb.close()
print(f'  17 disagreement 의 final consensus 적용: {len(consensus_data)} 케이스')

# disagreement 17 이 ADD15-balanced 안에 있는 케이스만 적용 (overlap 체크)
balanced_ids = {r['id'] for r in records}
in_balanced = [rid for rid in consensus_data if rid in balanced_ids]
print(f'  그 중 v5b-balanced 100 안: {len(in_balanced)}건')

# 합의 적용: 양 expert 가 final_consensus 라벨에 동의한 것으로 간주
# (단 final_consensus 가 uncertain 이면 gold 에서 제외)
records_with_consensus = []
for r in records:
    new_r = dict(r)
    if r['id'] in consensus_data:
        consensus = consensus_data[r['id']]
        if consensus in ('R1', 'not_R1'):
            new_r['sdh'] = consensus
            new_r['nkt'] = consensus
            new_r['consensus_applied'] = True
        else:
            new_r['consensus_applied'] = 'uncertain'
    records_with_consensus.append(new_r)

# 재평가
def kappa_v2(records, rule_key, label_key, name):
    pairs = []
    for r in records:
        rl = r[rule_key]; el = r[label_key]
        if rl in ('R1','not_R1') and el in ('R1','not_R1'):
            pairs.append((rl, el))
    if not pairs:
        print(f'  {name}: no valid pairs'); return
    yr, ye = zip(*pairs)
    yt = [['R1','not_R1'].index(x) for x in yr]
    yp = [['R1','not_R1'].index(x) for x in ye]
    raw = sum(1 for a,b in zip(yt,yp) if a==b)/len(yt)
    k = cohen_kappa_score(yt, yp)
    cm = confusion_matrix(yr, ye, labels=['R1','not_R1'])
    print(f'  {name}  n={len(pairs)}  κ={k:.3f}  raw={raw*100:.1f}%')
    print(f'    rule_R1→exp_R1={cm[0,0]} rule_R1→exp_not={cm[0,1]} | rule_not→exp_R1={cm[1,0]} rule_not→exp_not={cm[1,1]}')

print(f'\n=== After consensus applied — v5b rule vs expert ===')
kappa_v2(records_with_consensus, 'rule', 'sdh', 'SDH (consensus 반영) vs v5b')
kappa_v2(records_with_consensus, 'rule', 'nkt', 'NKT (consensus 반영) vs v5b')

# new gold
print(f'\n=== After consensus — gold-subset ===')
gold2 = [r for r in records_with_consensus if r['sdh'] == r['nkt'] and r['sdh'] in ('R1','not_R1')]
gold2_match = sum(1 for r in gold2 if r['rule'] == r['sdh'])
print(f'  gold n={len(gold2)}, rule match={gold2_match}/{len(gold2)} = {100*gold2_match/len(gold2):.1f}%')

# v5b R1 vs not_R1 breakdown
gold2_r1 = [r for r in gold2 if r['sdh'] == 'R1']
gold2_not = [r for r in gold2 if r['sdh'] == 'not_R1']
gold2_r1_m = sum(1 for r in gold2_r1 if r['rule'] == 'R1')
gold2_not_m = sum(1 for r in gold2_not if r['rule'] == 'not_R1')
print(f'  gold-R1     {gold2_r1_m}/{len(gold2_r1)} = {100*gold2_r1_m/len(gold2_r1) if gold2_r1 else 0:.1f}%')
print(f'  gold-not_R1 {gold2_not_m}/{len(gold2_not)} = {100*gold2_not_m/len(gold2_not) if gold2_not else 0:.1f}%')

# save records
import csv
with open('_v5b_final_records.csv', 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.DictWriter(f, fieldnames=['id','sample_no','origin','sdh','nkt','rule','consensus_applied'])
    w.writeheader()
    for r in records_with_consensus:
        w.writerow({k: r.get(k, '') for k in ['id','sample_no','origin','sdh','nkt','rule','consensus_applied']})
print(f'\n→ _v5b_final_records.csv')
