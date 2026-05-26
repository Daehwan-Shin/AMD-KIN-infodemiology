"""
전문가 2명(NKT, SDH) human label vs LLM consensus(model_final_v3) 평가.
- raw agreement, Cohen's kappa (4-class & binary)
- 전문가 간 일치도 (inter-rater)
- 불일치 케이스 표
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from collections import Counter, defaultdict

LABELS = ['Include_strict_AMD', 'Review_needed', 'Exclude_non_AMD', 'Unclear']


def cohens_kappa(pairs):
    n = len(pairs)
    if n == 0:
        return 0.0
    cats = sorted(set([a for a, b in pairs] + [b for a, b in pairs]))
    po = sum(1 for a, b in pairs if a == b) / n
    pe = sum(
        (sum(1 for a, _ in pairs if a == c) / n) *
        (sum(1 for _, b in pairs if b == c) / n)
        for c in cats
    )
    if pe == 1.0:
        return 1.0
    return (po - pe) / (1 - pe)


def read_review(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['review']
    hdr = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    rows = {}
    for r in range(2, ws.max_row + 1):
        no = ws.cell(row=r, column=hdr['No']).value
        if no is None:
            continue
        rows[int(no)] = {
            'id': ws.cell(row=r, column=hdr['id']).value,
            'label': (ws.cell(row=r, column=hdr['human_label']).value or '').strip(),
            'conf': ws.cell(row=r, column=hdr['human_confidence']).value,
        }
    return rows


def read_answerkey(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['answer_key']
    hdr = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    rows = {}
    for r in range(2, ws.max_row + 1):
        no = ws.cell(row=r, column=hdr['No']).value
        if no is None:
            continue
        rows[int(no)] = {
            'id': ws.cell(row=r, column=hdr['id']).value,
            'codex': ws.cell(row=r, column=hdr['codex']).value,
            'opus': ws.cell(row=r, column=hdr['opus']).value,
            'gemini': ws.cell(row=r, column=hdr['gemini']).value,
            'model_final': ws.cell(row=r, column=hdr['model_final_v3']).value,
        }
    return rows


answer = read_answerkey('AMD_human_validation_200_v3.xlsx')
nkt = read_review('AMD_human_validation_200_v3_NKT.xlsx')
sdh = read_review('AMD_human_validation_200_v3_SDH.xlsx')

print(f'answer_key: {len(answer)}건')
print(f'NKT 라벨링: {sum(1 for v in nkt.values() if v["label"])}/{len(nkt)}건 완료')
print(f'SDH 라벨링: {sum(1 for v in sdh.values() if v["label"])}/{len(sdh)}건 완료')

# id 정합성 체크
mismatch = [no for no in answer if answer[no]['id'] != nkt.get(no, {}).get('id')]
if mismatch:
    print(f'⚠ NKT id 불일치 No: {mismatch[:10]}')
mismatch2 = [no for no in answer if answer[no]['id'] != sdh.get(no, {}).get('id')]
if mismatch2:
    print(f'⚠ SDH id 불일치 No: {mismatch2[:10]}')


def binarize(lbl):
    return 'Include' if lbl == 'Include_strict_AMD' else 'NotInclude'


def report(name, human):
    print(f'\n{"="*64}')
    print(f'=== {name} vs LLM consensus (model_final_v3) ===')
    print(f'{"="*64}')
    valid = [no for no in answer if human.get(no, {}).get('label')]
    n = len(valid)
    if n == 0:
        print('  (라벨 없음 — 건너뜀)')
        return None

    pairs4 = [(human[no]['label'], answer[no]['model_final']) for no in valid]
    agree4 = sum(1 for a, b in pairs4 if a == b)
    pairsB = [(binarize(a), binarize(b)) for a, b in pairs4]
    agreeB = sum(1 for a, b in pairsB if a == b)

    print(f'  평가 건수            : {n}')
    print(f'  4-class raw agreement: {agree4}/{n} = {agree4/n*100:.1f}%')
    print(f'  4-class Cohen κ      : {cohens_kappa(pairs4):.3f}')
    print(f'  Binary(Include vs 그외) raw: {agreeB}/{n} = {agreeB/n*100:.1f}%')
    print(f'  Binary Cohen κ       : {cohens_kappa(pairsB):.3f}')

    # confusion matrix (human row, model col), binary
    cm = defaultdict(Counter)
    for h, m in pairsB:
        cm[h][m] += 1
    print(f'\n  [Binary confusion]  model→  Include  NotInclude')
    for h in ['Include', 'NotInclude']:
        print(f'    human={h:<11} {cm[h]["Include"]:>7} {cm[h]["NotInclude"]:>11}')

    # model_final이 Include인 것 중 human도 Include 비율 (PPV-like)
    model_inc = [no for no in valid if answer[no]['model_final'] == 'Include_strict_AMD']
    model_exc = [no for no in valid if answer[no]['model_final'] == 'Exclude_non_AMD']
    if model_inc:
        hit = sum(1 for no in model_inc if human[no]['label'] == 'Include_strict_AMD')
        print(f'\n  모델 Include {len(model_inc)}건 중 전문가도 Include: {hit} ({hit/len(model_inc)*100:.1f}%)')
    if model_exc:
        hit = sum(1 for no in model_exc if human[no]['label'] == 'Exclude_non_AMD')
        print(f'  모델 Exclude {len(model_exc)}건 중 전문가도 Exclude: {hit} ({hit/len(model_exc)*100:.1f}%)')

    # human 라벨 분포
    print(f'\n  {name} 라벨 분포: ', dict(Counter(human[no]["label"] for no in valid)))
    return {'pairs4': pairs4, 'valid': valid}


r_nkt = report('NKT', nkt)
r_sdh = report('SDH', sdh)

# === 전문가 간 일치도 ===
print(f'\n{"="*64}')
print(f'=== 전문가 간 일치도 (NKT vs SDH) ===')
print(f'{"="*64}')
both = [no for no in answer if nkt.get(no, {}).get('label') and sdh.get(no, {}).get('label')]
if both:
    pp = [(nkt[no]['label'], sdh[no]['label']) for no in both]
    ag = sum(1 for a, b in pp if a == b)
    ppB = [(binarize(a), binarize(b)) for a, b in pp]
    agB = sum(1 for a, b in ppB if a == b)
    print(f'  공통 평가 건수       : {len(both)}')
    print(f'  4-class raw agreement: {ag}/{len(both)} = {ag/len(both)*100:.1f}%')
    print(f'  4-class Cohen κ      : {cohens_kappa(pp):.3f}')
    print(f'  Binary raw agreement : {agB}/{len(both)} = {agB/len(both)*100:.1f}%')
    print(f'  Binary Cohen κ       : {cohens_kappa(ppB):.3f}')

    # 두 전문가 합의(동일 라벨) vs 모델
    consensus_no = [no for no in both if nkt[no]['label'] == sdh[no]['label']]
    cons_vs_model = sum(1 for no in consensus_no
                        if nkt[no]['label'] == answer[no]['model_final'])
    print(f'\n  두 전문가 합의 {len(consensus_no)}건 중 모델과도 일치: '
          f'{cons_vs_model} ({cons_vs_model/len(consensus_no)*100:.1f}%)')

# === 불일치 케이스 (모델 vs 두 전문가 모두 다른 것) ===
print(f'\n{"="*64}')
print(f'=== 모델이 두 전문가와 모두 불일치한 케이스 ===')
print(f'{"="*64}')
hard = []
for no in both:
    m = answer[no]['model_final']
    h1 = nkt[no]['label']
    h2 = sdh[no]['label']
    if h1 == h2 and h1 != m:
        hard.append((no, answer[no]['id'], m, h1,
                      answer[no]['codex'], answer[no]['opus'], answer[no]['gemini']))
print(f'  총 {len(hard)}건 (두 전문가가 같은 라벨인데 모델만 다름 = 모델 오류 의심)')
print(f'  {"No":>4} {"id":<9} {"model":<20} {"두전문가":<20} {"codex/opus/gemini"}')
for no, sid, m, h, c, o, g in sorted(hard)[:40]:
    print(f'  {no:>4} {str(sid):<9} {m:<20} {h:<20} {c}/{o}/{g}')

# 저장
with open('_stage3_eval_disagreements.txt', 'w', encoding='utf-8') as f:
    f.write(f'모델 vs 두 전문가 합의 불일치: {len(hard)}건\n\n')
    for no, sid, m, h, c, o, g in sorted(hard):
        f.write(f'No{no} | {sid} | model={m} | experts={h} | C={c} O={o} G={g}\n')
print(f'\n→ _stage3_eval_disagreements.txt 저장')
